import json
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from report_delivery.core import plan, process, verify, check_manifest
from report_delivery.delivery import deliver

HEADERS = ['统计日期', '视频ID', '视频标题', '抖音号ID', '视频发布时间',
           '累计播放次数', '累计直接成交金额', '累计直接成交券数',
           '累计直接核销金额', '累计直接核销券数', '新增播放次数']
EXPECTED = ['视频ID', '视频发布时间', '抖音号ID', '发布次数', '累计播放次数',
            '累计直接成交金额', '累计直接成交券数', '累计直接核销金额', '累计直接核销券数']
ROWS = [
    ['2026-01-01～2026-01-03', '900000000000000001', '合成示例甲', '000123', '2026/01/01 09:00:00', 100, '1,234.50', 2, 100, 1, 10],
    ['2026-01-01～2026-01-03', '900000000000000002', '合成示例乙', 'demo_account', '2026/01/02 10:00:00', 0, 0, 0, 0, 0, 0],
    ['2026-01-01～2026-01-03', '900000000000000001', '合成示例甲重复行', '000123', '2026/01/03 11:00:00', 200, 10, 1, 0, 0, 20],
]


class MemoryStorage:
    """External storage substitute: only the remote boundary is replaced."""
    def __init__(self, stale=False, uncertain=False):
        self.files = {}
        self.stale = stale
        self.uncertain = uncertain

    def stat(self, path):
        if path not in self.files:
            return None
        return {'path': path, 'size': len(self.files[path]), 'mtime': 1 if self.stale else 9999999999}

    def put(self, path, local):
        self.files[path] = Path(local).read_bytes()
        if self.uncertain:
            raise RuntimeError('Connection interrupted after request')


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'input.xlsx'
        self.write_source(ROWS)

    def write_source(self, rows, headers=HEADERS):
        wb = Workbook()
        wb.active.append(headers)
        for row in rows:
            wb.active.append(row)
        wb.save(self.source)
        wb.close()

    def prepared(self, demo=False):
        return process(self.source, self.root / 'run', '2026-01-01', '2026-01-03', demo=demo)

    def test_month_year_leap_boundaries(self):
        for when, start, end in [('2026-10-01', '2026-09-01', '2026-09-30'),
                                 ('2026-10-02', '2026-10-01', '2026-10-01'),
                                 ('2027-01-01', '2026-12-01', '2026-12-31'),
                                 ('2028-03-01', '2028-02-01', '2028-02-29')]:
            p = plan(datetime.fromisoformat(when+'T10:00:00+08:00'))
            self.assertEqual((p['start_date'], p['end_date']), (start, end))

    def test_preserves_order_duplicates_zero_rows_long_ids_and_style(self):
        record = self.prepared()
        self.assertEqual(record['status'], 'verified')
        self.assertEqual(record['rows'], 3)
        self.assertEqual(record['zero_rows_retained'], 1)
        wb = load_workbook(self.root / 'run' / record['output_file'])
        self.addCleanup(wb.close)
        ws = wb.active
        self.assertEqual([c.value for c in ws[1]], EXPECTED)
        self.assertEqual([ws.cell(i, 1).value for i in range(2, 5)], ['900000000000000001', '900000000000000002', '900000000000000001'])
        self.assertEqual(ws['C2'].value, '000123')
        self.assertEqual(ws['F2'].value, 1234.5)
        self.assertEqual([ws.cell(i, 4).value for i in range(2, 5)], [1, 1, 1])
        self.assertTrue(all(ws.column_dimensions[c].width == 18 for c in 'ABCDEFGHI'))
        self.assertEqual(ws.sheet_format.defaultRowHeight, 16.8)
        self.assertEqual(ws['D1'].font.name, '宋体-简')
        self.assertEqual(ws['A2'].font.name, 'Calibri')
        self.assertEqual(ws['A2'].number_format, 'General')
        self.assertEqual(ws['A2'].data_type, 's')

    def test_empty_missing_duplicate_headers_and_wrong_period_block_processing(self):
        cases = [([], HEADERS), (ROWS, HEADERS[1:]), (ROWS, HEADERS+['视频ID'])]
        for rows, headers in cases:
            self.write_source(rows, headers)
            with self.assertRaises(ValueError):
                self.prepared()
        bad = [list(row) for row in ROWS]
        bad[0][0] = '2025-12-01～2025-12-31'
        self.write_source(bad)
        with self.assertRaises(ValueError):
            self.prepared()

    def test_unsafe_numeric_ids_formula_and_out_of_range_date_block_processing(self):
        for column, value in [(1, 900000000000000001), (5, '=1+1'), (4, '2026/02/01 00:00:00')]:
            rows = [list(row) for row in ROWS]
            rows[0][column] = value
            self.write_source(rows)
            with self.assertRaises(ValueError):
                self.prepared()

    def test_verification_detects_data_order_or_font_changes(self):
        record = self.prepared()
        target = self.root / 'run' / record['output_file']
        original = target.read_bytes()
        for mutation in ['value', 'order', 'font', 'width']:
            target.write_bytes(original)
            wb = load_workbook(target)
            if mutation == 'value': wb.active['E2'] = 999
            if mutation == 'order': wb.active['A2'] = '900000000000000002'
            if mutation == 'font': wb.active['A2'].font = Font(name='Arial', size=11)
            if mutation == 'width': wb.active.column_dimensions['A'].width = 30
            wb.save(target)
            wb.close()
            with self.assertRaises(ValueError):
                verify(self.source, target, '2026-01-01', '2026-01-03')

    def test_reuse_checks_hashes_and_refuses_modified_data(self):
        first = self.prepared()
        self.assertTrue(self.prepared()['reused'])
        (self.root / 'run' / first['output_file']).write_bytes(b'changed')
        with self.assertRaises(ValueError):
            check_manifest(self.root / 'run' / 'manifest.json')

    def test_dry_run_does_not_contact_storage(self):
        self.prepared()
        result = deliver(self.root/'run/manifest.json', '/reports', None, execute=False)
        self.assertEqual(result['status'], 'planned')
        self.assertEqual(json.loads((self.root/'run/manifest.json').read_text())['status'], 'verified')

    def test_demo_is_blocked_from_real_upload(self):
        self.prepared(demo=True)
        with self.assertRaises(ValueError):
            deliver(self.root/'run/manifest.json', '/reports', MemoryStorage(), execute=True)

    def test_upload_records_readback_then_reuses_completion(self):
        record = self.prepared()
        storage = MemoryStorage()
        path = self.root/'run/manifest.json'
        result = deliver(path, '/reports', storage, execute=True)
        self.assertEqual(result['status'], 'uploaded')
        self.assertEqual(storage.files['/reports/202601短视频数据源.xlsx'], (path.parent/record['output_file']).read_bytes())
        self.assertFalse(result['upload_verification']['remote_hash_checked'])
        self.assertTrue(deliver(path, '/reports', None, execute=True)['reused'])
        with self.assertRaises(ValueError):
            deliver(path, '/other-folder', None, execute=True)

    def test_stale_readback_or_uncertain_upload_never_marks_completed(self):
        self.prepared()
        path = self.root/'run/manifest.json'
        with self.assertRaises(RuntimeError):
            deliver(path, '/reports', MemoryStorage(stale=True), execute=True)
        self.assertEqual(json.loads(path.read_text())['status'], 'verified')
        with self.assertRaises(RuntimeError):
            deliver(path, '/reports', MemoryStorage(), execute=True)

    def test_manifest_cannot_escape_run_directory(self):
        self.prepared()
        path = self.root/'run/manifest.json'
        record = json.loads(path.read_text())
        record['source_file'] = '../input.xlsx'
        path.write_text(json.dumps(record))
        with self.assertRaises(ValueError):
            check_manifest(path)

    def test_completed_upload_is_bound_to_server_as_well_as_folder(self):
        self.prepared()
        path = self.root/'run/manifest.json'
        result = deliver(path, '/reports', MemoryStorage(), execute=True,
                         endpoint='https://nas-a.example.com')
        self.assertEqual(result['upload_verification']['endpoint'], 'https://nas-a.example.com')
        with self.assertRaises(ValueError):
            deliver(path, '/reports', endpoint='https://nas-b.example.com')

    def test_verification_detects_font_color_or_underline_changes(self):
        record = self.prepared()
        target = self.root/'run'/record['output_file']
        wb = load_workbook(target)
        wb.active['A2'].font = Font(name='Calibri', size=11, color='FF0000', underline='single')
        wb.save(target)
        wb.close()
        with self.assertRaises(ValueError):
            verify(self.source, target, '2026-01-01', '2026-01-03')


if __name__ == '__main__':
    unittest.main()
