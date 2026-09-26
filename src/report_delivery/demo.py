"""An entirely synthetic fixture and local delivery demonstration; no network."""
import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .core import process, save_json, sha256

SOURCE_HEADERS = ['统计日期', '视频ID', '视频标题', '抖音号ID', '视频发布时间',
    '累计播放次数', '累计直接成交金额', '累计直接成交券数', '累计直接核销金额', '累计直接核销券数', '新增播放次数']
SYNTHETIC_ROWS = [
    ['2026-01-01～2026-01-03', '900000000000000001', '合成示例甲', '000123', '2026/01/01 09:00:00', 100, '1,234.50', 2, 100, 1, 10],
    ['2026-01-01～2026-01-03', '900000000000000002', '合成示例乙', 'demo_account', '2026/01/02 10:00:00', 0, 0, 0, 0, 0, 0],
    ['2026-01-01～2026-01-03', '900000000000000001', '合成示例甲重复行', '000123', '2026/01/03 11:00:00', 200, 10, 1, 0, 0, 20],
]


def run_demo(directory):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    source = directory/'synthetic-input.xlsx'
    if source.exists():
        wb = load_workbook(source, read_only=True, data_only=False)
        try:
            values = [list(row) for row in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
        if values != [SOURCE_HEADERS, *SYNTHETIC_ROWS]:
            raise ValueError('Demo directory contains a different input; choose a new directory')
    else:
        wb = Workbook()
        wb.properties.creator = 'Synthetic demo'
        wb.active.append(SOURCE_HEADERS)
        for row in SYNTHETIC_ROWS:
            wb.active.append(row)
        wb.save(source)
        wb.close()
    record = process(source, directory/'run', '2026-01-01', '2026-01-03', demo=True)
    archive = directory/'local-delivery'
    archive.mkdir(exist_ok=True)
    target = archive/record['output_file']
    if target.exists() and sha256(target) != record['output_sha256']:
        raise ValueError('Existing demo delivery differs; refusing to overwrite')
    if not target.exists():
        shutil.copyfile(directory/'run'/record['output_file'], target)
    if sha256(target) != record['output_sha256']:
        raise RuntimeError('Local demo copy failed hash verification')
    result = {'status': 'demo_completed', 'rows': record['rows'],
              'zero_rows_retained': record['zero_rows_retained'], 'network_accessed': False,
              'delivery_method': 'local copy with SHA256 verification',
              'manifest': 'run/manifest.json', 'delivered_file': 'local-delivery/'+target.name}
    save_json(directory/'demo-result.json', result)
    return result
