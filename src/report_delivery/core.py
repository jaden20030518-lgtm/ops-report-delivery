"""Deterministic report transformation, independent re-read and file evidence."""
import hashlib
import json
import os
import re
import shutil
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from itertools import zip_longest
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

HEADERS = ['视频ID', '视频发布时间', '抖音号ID', '发布次数', '累计播放次数',
           '累计直接成交金额', '累计直接成交券数', '累计直接核销金额', '累计直接核销券数']
BEIJING = ZoneInfo('Asia/Shanghai')


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def save_json(path, value):
    path = Path(path)
    temporary = path.with_suffix(path.suffix + '.tmp')
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
    os.replace(temporary, path)


@contextmanager
def run_lock(directory):
    lock = Path(directory) / '.run-lock'
    try:
        lock.mkdir()
    except FileExistsError:
        raise RuntimeError('Run is locked; inspect the previous process before recovery') from None
    try:
        yield
    finally:
        lock.rmdir()


def plan(now=None):
    now = now or datetime.now(BEIJING)
    if now.tzinfo is None:
        raise ValueError('Timestamp must include a timezone')
    end = now.astimezone(BEIJING).date() - timedelta(days=1)
    return {'start_date': end.replace(day=1).isoformat(), 'end_date': end.isoformat(),
            'filename': end.strftime('%Y%m') + '短视频数据源.xlsx', 'timezone': 'Asia/Shanghai'}


def window(start, end):
    start, end = date.fromisoformat(start), date.fromisoformat(end)
    if start > end or start != end.replace(day=1):
        raise ValueError('Range must start on the first day of the ending month')
    return start, end


def calendar_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.fromisoformat(str(value).replace('/', '-')).date()
    except ValueError:
        raise ValueError('Invalid publication date') from None


def read_source(source, start, end):
    """Read and validate the source independently of the output workbook."""
    first, last = window(start, end)
    wb = load_workbook(source, read_only=True, data_only=False)
    try:
        if len(wb.worksheets) != 1:
            raise ValueError('Source must contain exactly one worksheet')
        sheet = wb.active
        sheet.reset_dimensions()
        iterator = sheet.iter_rows()
        headers = [cell.value for cell in next(iterator, ())]
        required = [h for h in HEADERS if h != '发布次数'] + ['视频标题', '统计日期']
        for header in required:
            if headers.count(header) != 1:
                raise ValueError(f'Missing or duplicate header: {header}')
        positions = {h: headers.index(h) for h in required}
        rows = []
        for cells in iterator:
            values = [cell.value for cell in cells]
            if not any(value is not None for value in values):
                continue
            if any(cell.data_type in ('f', 'e') for cell in cells):
                raise ValueError('Source formulas and Excel error cells are not supported')
            values += [None] * max(0, len(headers) - len(values))
            period = re.findall(r'\d{4}[-/]\d{1,2}[-/]\d{1,2}', str(values[positions['统计日期']]))
            dates = [calendar_date(value) for value in period]
            if dates != [first, last] and not (first == last and dates == [first]):
                raise ValueError('Statistical period does not match the requested range')
            publication = values[positions['视频发布时间']]
            if not first <= calendar_date(publication) <= last:
                raise ValueError('Publication date is outside the requested range')
            row = []
            for header in HEADERS:
                value = 1 if header == '发布次数' else values[positions[header]]
                if header in ('视频ID', '抖音号ID'):
                    if isinstance(value, (float, bool)) or value is None or value == '':
                        raise ValueError('Identifiers must be present and lossless')
                    if isinstance(value, int) and len(str(abs(value))) > 15:
                        raise ValueError('Long identifiers must be stored as source text')
                    value = str(value)
                elif header == '视频发布时间' and isinstance(value, (datetime, date)):
                    value = value.strftime('%Y/%m/%d %H:%M:%S')
                elif header in HEADERS[4:] and value is not None:
                    try:
                        value = Decimal(str(value).replace(',', ''))
                    except InvalidOperation:
                        raise ValueError('Metric is not numeric') from None
                    if not value.is_finite():
                        raise ValueError('Metric must be finite')
                row.append(value)
            rows.append(row)
        if not rows:
            raise ValueError('Empty report is not eligible for delivery')
        return rows
    finally:
        wb.close()


def verify(source, output, start, end):
    expected = read_source(source, start, end)
    wb = load_workbook(output, data_only=False)
    try:
        if wb.sheetnames != ['Sheet1']:
            raise ValueError('Output must contain only Sheet1')
        ws = wb.active
        rows = ws.iter_rows()
        if [cell.value for cell in next(rows, ())] != HEADERS:
            raise ValueError('Output header names or order changed')
        if ws.sheet_format.defaultRowHeight != 16.8:
            raise ValueError('Default row height changed')
        if any(ws.column_dimensions[c].width != 18 for c in 'ABCDEFGHI'):
            raise ValueError('Column width changed')
        if ws.merged_cells or ws.tables or ws.freeze_panes or ws.auto_filter.ref:
            raise ValueError('Unexpected workbook decoration')
        if any(dim.height is not None or dim.has_style for dim in ws.row_dimensions.values()):
            raise ValueError('Unexpected per-row height or style')
        for row in ws.iter_rows():
            for cell in row:
                font = '宋体-简' if cell.coordinate == 'D1' else 'Calibri'
                if cell.font != Font(name=font, size=11):
                    raise ValueError('Font changed')
                if cell.number_format != 'General' or cell.fill.patternType is not None:
                    raise ValueError('Cell format changed')
        for original, cells in zip_longest(expected, rows):
            if original is None or cells is None or len(cells) != 9:
                raise ValueError('Output row or column count changed')
            for column, (value, cell) in enumerate(zip(original, cells)):
                if column in (0, 2):
                    valid = cell.data_type == 's' and cell.value == value
                elif column == 1:
                    valid = cell.value == value and cell.data_type == 's'
                elif value is None:
                    valid = cell.value is None
                else:
                    valid = cell.data_type == 'n' and Decimal(str(cell.value)) == Decimal(str(value))
                if not valid:
                    raise ValueError(f'Output differs from source at {cell.coordinate}')
        return {'rows': len(expected), 'columns': 9,
                'zero_rows_retained': sum(all(v == 0 for v in row[4:]) for row in expected),
                'all_values_verified': True, 'format_verified': True}
    finally:
        wb.close()


def check_manifest(path):
    path = Path(path)
    record = json.loads(path.read_text(encoding='utf-8'))
    if record.get('status') not in ('verified', 'uploaded') or record.get('rows', 0) < 1:
        raise ValueError('Manifest has not passed verification')
    _, end = window(record['start_date'], record['end_date'])
    if record['output_file'] != end.strftime('%Y%m') + '短视频数据源.xlsx':
        raise ValueError('Output filename does not match ending month')
    for field in ('source', 'output'):
        name = record[field+'_file']
        if not isinstance(name, str) or Path(name).name != name or name in ('.', '..'):
            raise ValueError('Manifest file paths must stay in the run directory')
        target = path.parent/name
        if target.is_symlink() or sha256(target) != record[field+'_sha256']:
            raise ValueError('Verified file was modified')
    if (path.parent/record['output_file']).stat().st_size != record['output_bytes']:
        raise ValueError('Output size changed')
    return record


def process(source, directory, start, end, *, demo=False):
    source, directory = Path(source), Path(directory)
    data = read_source(source, start, end)
    directory.mkdir(parents=True, exist_ok=True)
    with run_lock(directory):
        manifest = directory/'manifest.json'
        if manifest.exists():
            old = check_manifest(manifest)
            if (old['source_sha256'] != sha256(source) or old['start_date'] != start
                    or old['end_date'] != end or old['demo'] != demo):
                raise ValueError('Run directory belongs to a different input or plan')
            verify(directory/old['source_file'], directory/old['output_file'], start, end)
            return {**old, 'reused': True}
        raw = directory/'source.xlsx'
        output = directory/(end[:7].replace('-', '')+'短视频数据源.xlsx')
        if raw.exists() or output.exists():
            raise ValueError('Incomplete previous run exists; preserve it for inspection')
        record = {'schema_version': 1, 'status': 'processing', 'demo': demo,
                  'ready_for_upload': False, 'start_date': start, 'end_date': end,
                  'source_file': raw.name, 'output_file': output.name,
                  'source_sha256': sha256(source)}
        save_json(manifest, record)
        try:
            shutil.copyfile(source, raw)
            if sha256(raw) != record['source_sha256']:
                raise ValueError('Source changed during backup')
            wb = Workbook()
            ws = wb.active
            ws.title = 'Sheet1'
            ws.sheet_format.defaultRowHeight = 16.8
            for column in 'ABCDEFGHI':
                ws.column_dimensions[column].width = 18
            ws.append(HEADERS)
            for row in data:
                ws.append(row)
            normal = Font(name='Calibri', size=11)
            for row in ws.iter_rows():
                for cell in row:
                    cell.font = normal
                    cell.number_format = 'General'
            ws['D1'].font = Font(name='宋体-简', size=11)
            wb.save(output)
            wb.close()
            checks = verify(raw, output, start, end)
            record.update(checks, status='verified', ready_for_upload=not demo,
                          output_sha256=sha256(output), output_bytes=output.stat().st_size)
            save_json(manifest, record)
            return record
        except Exception:
            record.update(status='failed', ready_for_upload=False)
            save_json(manifest, record)
            raise
